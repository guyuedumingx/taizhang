"""
员工批量导入 (2026-10-16 决策 2) 测试

覆盖:
  - 4 列模板 (EHR号/姓名/组别/初始密码) 正常导入
  - 初始密码强度: 缺大写/小写/数字/特殊字符/长度不足 -> 整行失败, 明细列出缺哪几项
  - 组别按名字匹配, 匹配不到 -> 整行失败 (不是静默留空)
  - EHR 格式错 / 重复 -> 整行失败
  - 缺列 -> 400 并列出全部缺失列
  - Excel 把 7 位 EHR 读成 "1234567.0" 的兼容
  - 模板下载端点 4 列齐全
  - 普通员工无权批量建号 (PRD §13)

隔离: 根 tests/conftest.py 已推广函数级内存库 (重构方案 Phase 1.1),
本文件不自带 shim, 直接用根固件。
"""
import io
import random

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app import models
from app.core.config import settings
from app.core.security import create_access_token
from app.main import app

client = TestClient(app)
API_USERS = f"{settings.API_V1_STR}/users"
API_PORTRAIT = f"{settings.API_V1_STR}/portrait"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="function")
def team(db: Session):
    t = models.Team(name="导入测试组", department="测试部")
    db.add(t)
    db.commit()
    db.refresh(t)
    return t


def _ehr(db: Session) -> str:
    while True:
        e = str(random.randint(3000000, 8999999))
        if not db.query(models.User).filter(models.User.ehr_id == e).first():
            return e


def _row(db: Session, team, **over) -> dict:
    row = {
        "EHR号": _ehr(db),
        "姓名": "导入员工",
        "组别": team.name,
        "初始密码": "Ab3#efGh9",
    }
    row.update(over)
    return row


def _xlsx(rows: list[dict]) -> bytes:
    import pandas as pd

    buf = io.BytesIO()
    pd.DataFrame(rows).to_excel(buf, index=False)
    return buf.getvalue()


def _post(rows, headers):
    return client.post(
        f"{API_USERS}/import",
        headers=headers,
        files={"file": ("users.xlsx", _xlsx(rows), XLSX_MIME)},
    )


# ==================================================================== 测试
def test_import_success(db: Session, admin_token_headers, team):
    """4 列齐全且合法 → 导入成功, 组别落到正确的组"""
    ehr = _ehr(db)
    resp = _post([_row(db, team, EHR号=ehr, 姓名="张三")], admin_token_headers)

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert body["success_count"] == 1, body
    assert body["failed_count"] == 0, body

    u = db.query(models.User).filter(models.User.ehr_id == ehr).first()
    assert u is not None, "导入的 EHR 在库里查不到"
    assert u.name == "张三"
    assert u.team_id == team.id, f"组别没落上, team_id={u.team_id}"
    assert u.is_superuser is False, "导入的普通员工不该是超管"


def test_weak_password_rejected(db: Session, admin_token_headers, team):
    """弱密码整行失败, 且明细要说清缺哪几项 —— 不能静默接受弱密码"""
    resp = _post([_row(db, team, 初始密码="abc12345")], admin_token_headers)
    body = resp.json()

    assert body["success_count"] == 0, body
    assert body["failed_count"] == 1, body
    reason = body["failed_users"][0]["reason"]
    assert "大写字母" in reason, reason
    assert "特殊字符" in reason, reason


def test_password_too_short_rejected(db: Session, admin_token_headers, team):
    """四类字符都齐但不足 8 位, 仍要失败"""
    resp = _post([_row(db, team, 初始密码="aB3#efG")], admin_token_headers)
    body = resp.json()
    assert body["success_count"] == 0, body
    assert "8" in body["failed_users"][0]["reason"]


def test_unknown_team_rejected(db: Session, admin_token_headers, team):
    """组别不存在必须整行失败。

    原实现是 team_id 匹配不到就留 team_id=None —— 员工成了无组孤儿,
    而本项目按组判组长, 错得没有任何提示 (Rule 12)。
    """
    resp = _post([_row(db, team, 组别="根本不存在的组")], admin_token_headers)
    body = resp.json()

    assert body["success_count"] == 0, body
    reason = body["failed_users"][0]["reason"]
    assert "组别不存在" in reason, reason
    assert "根本不存在的组" in reason, f"明细应指名是哪个组: {reason}"


def test_bad_ehr_format_rejected(db: Session, admin_token_headers, team):
    """EHR 必须是 7 位数字"""
    resp = _post([_row(db, team, EHR号="12345")], admin_token_headers)
    body = resp.json()
    assert body["success_count"] == 0, body
    assert "EHR号格式错误" in body["failed_users"][0]["reason"]


def test_ehr_read_as_float(db: Session, admin_token_headers, team):
    """Excel 把 7 位 EHR 当数字存时会读成 '1234567.0', 必须还原, 否则判格式错"""
    ehr = _ehr(db)
    resp = _post([_row(db, team, EHR号=f"{ehr}.0")], admin_token_headers)
    body = resp.json()

    assert body["success_count"] == 1, f"带 .0 的 EHR 应被接受: {body}"
    u = db.query(models.User).filter(models.User.ehr_id == ehr).first()
    assert u is not None, f"库里应是 7 位纯数字, 实际查不到 {ehr}"


def test_duplicate_ehr_rejected(db: Session, admin_token_headers, team):
    """EHR 重复整行失败, 明细点名已存在的是谁"""
    ehr = _ehr(db)
    first = _post([_row(db, team, EHR号=ehr, 姓名="原有人")], admin_token_headers)
    assert first.json()["success_count"] == 1, first.text

    resp = _post([_row(db, team, EHR号=ehr, 姓名="新人")], admin_token_headers)
    body = resp.json()
    assert body["success_count"] == 0, body
    reason = body["failed_users"][0]["reason"]
    assert "EHR号已存在" in reason, reason
    assert "原有人" in reason, f"明细应指出占用的人: {reason}"


def test_missing_columns_lists_all(db: Session, admin_token_headers):
    """缺列 -> 400, 且要一次列全缺了哪几列, 不让管理员来回试"""
    import pandas as pd

    buf = io.BytesIO()
    pd.DataFrame([{"EHR号": "1234567", "姓名": "张三"}]).to_excel(buf, index=False)
    resp = client.post(
        f"{API_USERS}/import",
        headers=admin_token_headers,
        files={"file": ("u.xlsx", buf.getvalue(), XLSX_MIME)},
    )

    assert resp.status_code == 400, resp.text
    detail = resp.json()["detail"]
    assert "组别" in detail, detail
    assert "初始密码" in detail, detail


def test_partial_failure_is_visible(db: Session, admin_token_headers, team):
    """部分成功部分失败时, 失败必须逐行列出来 (Rule 12: 跳过数量和原因必须可见)"""
    good = _ehr(db)
    resp = _post(
        [
            _row(db, team, EHR号=good, 姓名="好的那行"),
            _row(db, team, 初始密码="weak"),
            _row(db, team, 组别="查无此组"),
        ],
        admin_token_headers,
    )
    body = resp.json()

    assert body["success_count"] == 1, body
    assert body["failed_count"] == 2, body
    reasons = " | ".join(u["reason"] for u in body["failed_users"])
    assert "初始密码不合规" in reasons, reasons
    assert "组别不存在" in reasons, reasons


def test_import_denied_for_normal_user(db: Session, normal_token_headers, team):
    """普通员工不能批量建号 (PRD §13)"""
    resp = _post([_row(db, team)], normal_token_headers)
    assert resp.status_code == 403, resp.text


# ==================================================================== 模板
def test_user_import_template_has_4_columns():
    """模板必须是决策 2 定的 4 列, 且表头与导入端点认的一致"""
    resp = client.get(f"{API_PORTRAIT}/templates/excel/user-import")
    assert resp.status_code == 200, resp.text

    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(resp.content)).active
    headers = [c.value for c in ws[1]]

    assert headers == ["EHR号", "姓名", "组别", "初始密码"], headers
